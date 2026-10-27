"""Rebuild the ZONE_MASTER, ZONE_LINEAGE_MASTER, ZONE_OWNER_CARD_MASTER and ZONE_OWNER_FT_MASTER sheets of
ZONE_MASTER.xlsx
from ZONE_BOUNDARY_MASTER.kml.

Sources:
- KML (source of truth): zone list and order, polygons, and every per-zone value as <h1>field: value</h1>
  lines in the balloon text of the zone's own style, which Google Earth shows as its "Description"
  (see DESCRIPTION_FIELDS)
- layouts/CAMPUS_MASTER.md: company and campus columns

Runs validate_zone_master.py on the result and exits non-zero if it fails. See layouts/ZONE_LAYOUTS.md §6.

Usage (from the repo root):
    .venv/bin/python src/build_zone_master.py layouts/master_layout/ZONE_BOUNDARY_MASTER.kml
"""
import datetime as dt
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from shapely.geometry import Polygon

from kml_to_shp import NS, style_balloons, style_maps

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "layouts/master_layout/ZONE_MASTER.xlsx"
CAMPUS_MD = ROOT / "layouts/CAMPUS_MASTER.md"

NULL = "NULL"

# Per-zone values stored in each zone's Description, in this order.
DESCRIPTION_FIELDS = [
    "record_id", "zone_code", "production_unit_name", "production_unit_code", "production_unit_type",
    "zone_type", "zone_type_confidence", "primary_product_association", "product_association_confidence",
    "effective_from", "effective_to", "pit_available_from", "pit_available_to", "imagery_acquired_date",
]

COLUMNS = [
    "record_id", "zone_code", "company_name_en", "security_code", "campus_name", "campus_code",
    "campus_address", "production_unit_name", "production_unit_code", "production_unit_type",
    "zone_type", "zone_type_confidence", "primary_product_association", "product_association_confidence",
    "effective_from", "effective_to", "pit_available_from", "pit_available_to", "imagery_acquired_date",
    "record_created_at", "polygon",
]

# ZONE_OWNER_CARD_MASTER: current campus-division (CD) zones only, with ZONE_MASTER's column names.
CARD_SHEET = "ZONE_OWNER_CARD_MASTER"
CARD_COLUMNS = ["zone_code", "polygon"]

# ZONE_OWNER_FT_MASTER: every zone, latest version of each record_id (ended zones included), for the
# foot-traffic data owner (KT), with ZONE_MASTER's column names.
FT_SHEET = "ZONE_OWNER_FT_MASTER"
FT_COLUMNS = ["zone_code", "effective_from", "effective_to", "polygon"]


def is_cd(row):
    """Campus divisions are a level, not a zone type: zone_type is NULL and the zone_code keeps CD (SKH_ICH_CD_01)."""
    return row["zone_code"].split("_")[2] == "CD"


def read_kml(kml_path):
    """Return placemarks in document order: folder, name, description key/values, raw ring text."""
    doc = ET.parse(kml_path).getroot().find("k:Document", NS)
    balloons, maps = style_balloons(doc), style_maps(doc)
    style_owner = {}
    out = []
    for folder in doc.iter("{%s}Folder" % NS["k"]):
        folder_name = folder.find("k:name", NS).text
        for pm in folder.findall("k:Placemark", NS):
            name = pm.find("k:name", NS).text
            style_url = pm.find("k:styleUrl", NS).text.lstrip("#")
            if style_url in style_owner:
                sys.exit(f"{name}: shares its style with {style_owner[style_url]}; editing one zone's Description "
                         "in Google Earth would change both. Give each zone its own style.")
            style_owner[style_url] = name
            text = balloons.get(maps.get(style_url, "#" + style_url).lstrip("#")) or ""
            attrs = dict(re.findall(r"<h1>\s*([a-z_]+):\s*(.*?)\s*</h1>", text))
            missing = [f for f in DESCRIPTION_FIELDS if f not in attrs]
            if missing:
                sys.exit(f"{name}: Description is missing {missing}")
            if attrs["zone_code"] != name:
                sys.exit(f"{name}: Description zone_code is {attrs['zone_code']!r}")
            rings = pm.findall("k:Polygon/k:outerBoundaryIs/k:LinearRing/k:coordinates", NS)
            if len(rings) != 1 or pm.findall("k:Polygon/k:innerBoundaryIs", NS):
                sys.exit(f"{name}: unexpected geometry structure")
            out.append(dict(folder=folder_name, name=name, attrs=attrs, polygon=" ".join(rings[0].text.split())))
    total = len(list(doc.iter("{%s}Placemark" % NS["k"])))
    if len(out) != total:
        sys.exit(f"{total - len(out)} placemark(s) are not directly inside a campus folder")
    return out


def ring_polygon(raw):
    return Polygon([tuple(float(v) for v in tok.split(",")[:2]) for tok in raw.split()])


def read_campus_master(md_path):
    """Parse the campus table in CAMPUS_MASTER.md into {campus_code: row}."""
    lines = [ln for ln in md_path.read_text(encoding="utf-8").splitlines() if ln.startswith("|")]
    header = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = {}
    for ln in lines[2:]:
        row = dict(zip(header, (c.strip() for c in ln.strip("|").split("|"))))
        rows[row["campus_code"]] = row
    return rows


def build_rows(zones, campuses, created_at):
    rows, notes = [], []
    for z in zones:
        code = z["name"]
        campus_code = "_".join(code.split("_")[:2])
        if campus_code != z["folder"]:
            sys.exit(f"{code}: campus {campus_code} != KML folder {z['folder']}")
        campus = campuses[campus_code]
        row = {f: z["attrs"][f] for f in DESCRIPTION_FIELDS}
        row.update(
            company_name_en=campus["company_name_en"],
            security_code=campus["security_code"],
            campus_name=campus["campus_name"],
            campus_code=campus_code,
            campus_address=campus["campus_address"],
            record_created_at=created_at,
            polygon=z["polygon"],
        )
        rows.append(row)
    return rows, notes


# Lineage parent zone_type -> zone types it may become. Only EX parents may SPLIT (an unfinished remainder stays EX).
# CD rows (zone_type NULL) are never parents or children.
LINEAGE_CHILDREN = {"EX": {"EX", "OP", "OAE"}, "OP": {"OAE"}, "OAE": {"OP"}}


def build_lineage(rows):
    """TRANSFORM: a zone ends and one zone with the identical polygon starts the next day in the same campus
    (EX -> OP when operation starts, EX -> OAE when operation starts while construction continues in the same
    building, OP -> OAE when works start on top of an operating building, OAE -> OP when they end).
    SPLIT: several zones start the next day inside an EX polygon."""
    lineage, notes, seq = [], [], {}
    for parent in rows:
        ptype = parent["zone_type"]
        if ptype not in LINEAGE_CHILDREN or parent["effective_to"] == NULL:
            continue
        allowed = LINEAGE_CHILDREN[ptype]
        next_day = (dt.date.fromisoformat(parent["effective_to"]) + dt.timedelta(days=1)).isoformat()
        ppoly = ring_polygon(parent["polygon"])
        kids = [
            r for r in rows
            if r is not parent and r["campus_code"] == parent["campus_code"] and r["zone_type"] in allowed
            and r["effective_from"] == next_day
            and ppoly.buffer(1e-7).contains(ring_polygon(r["polygon"]))
        ]
        if len(kids) == 1 and ring_polygon(kids[0]["polygon"]).equals(ppoly):
            rel = "TRANSFORM"
        elif len(kids) > 1 and ptype == "EX":
            rel = "SPLIT"
        else:
            if ptype != "OP" or kids:  # an OP zone may simply close
                notes.append(f"{parent['zone_code']}: ends {parent['effective_to']} but no valid successor zone found")
            continue
        for kid in kids:
            short = parent["campus_code"].split("_")[1]
            seq[short] = seq.get(short, 0) + 1
            lineage.append(dict(
                lineage_id=f"LIN_{short}_{seq[short]:03d}",
                parent_record_id=parent["record_id"],
                child_record_id=kid["record_id"],
                relationship_type=rel,
                effective_at=next_day,
                pit_available_from=next_day,
                pit_available_to=NULL,
                _pair=f"{parent['zone_code']} -> {kid['zone_code']}",
            ))
    return lineage, notes


def write_zone_master(wb, rows):
    """Replace the ZONE_MASTER sheet, keeping header styling and column widths but dropping check_status(TEMP)."""
    old = wb["ZONE_MASTER"]
    old_header = [c.value for c in old[1]]
    offset = 1 if old_header[:1] == ["check_status(TEMP)"] else 0
    if old_header[offset:] != COLUMNS:
        sys.exit(f"ZONE_MASTER header changed unexpectedly: {old_header}")
    # On the first run the old columns B..V become A..U; later runs keep A..U.
    styles = [old.cell(row=1, column=i + 1 + offset) for i in range(len(COLUMNS))]
    widths = [old.column_dimensions[c.column_letter].width for c in styles]
    header_height = old.row_dimensions[1].height
    index = wb.sheetnames.index("ZONE_MASTER")
    wb.remove(old)

    ws = wb.create_sheet("ZONE_MASTER", index)
    for col, (name, src, width) in enumerate(zip(COLUMNS, styles, widths), start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill, cell.border = copy(src.font), copy(src.fill), copy(src.border)
        cell.alignment, cell.number_format = copy(src.alignment), src.number_format
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = header_height
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
    ws.freeze_panes = "A2"
    wb.active = index


def write_lineage(wb, lineage):
    ws = wb["ZONE_LINEAGE_MASTER"]
    header = [c.value for c in ws[1] if c.value]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in lineage:
        ws.append([r[c] for c in header])


def write_copy_sheet(wb, sheet, columns, rows, after):
    """Replace `sheet` (or create it right after `after`) with `columns` of `rows`, styled like ZONE_MASTER."""
    if sheet in wb.sheetnames:
        index = wb.sheetnames.index(sheet)
        wb.remove(wb[sheet])
    else:
        index = wb.sheetnames.index(after) + 1
    ws = wb.create_sheet(sheet, index)
    zm = wb["ZONE_MASTER"]
    src = {c.value: c for c in zm[1]}
    for col, name in enumerate(columns, start=1):
        s = src[name]
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill, cell.border = copy(s.font), copy(s.fill), copy(s.border)
        cell.alignment, cell.number_format = copy(s.alignment), s.number_format
        ws.column_dimensions[cell.column_letter].width = zm.column_dimensions[s.column_letter].width
    ws.row_dimensions[1].height = zm.row_dimensions[1].height
    for r in rows:
        ws.append([r[c] for c in columns])
    ws.freeze_panes = "A2"
    return len(rows)


def latest_versions(rows):
    """One row per record_id: the version with the latest pit_available_from, in ZONE_MASTER order."""
    latest = {}
    for r in rows:
        if r["record_id"] not in latest or r["pit_available_from"] > latest[r["record_id"]]["pit_available_from"]:
            latest[r["record_id"]] = r
    return [r for r in rows if latest[r["record_id"]] is r]


def main(kml_path):
    zones = read_kml(kml_path)
    campuses = read_campus_master(CAMPUS_MD)
    created_at = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M")

    rows, notes = build_rows(zones, campuses, created_at)
    lineage, lineage_notes = build_lineage(rows)

    wb = load_workbook(XLSX_PATH)
    write_zone_master(wb, rows)
    write_lineage(wb, lineage)
    n_card = write_copy_sheet(wb, CARD_SHEET, CARD_COLUMNS,
                              [r for r in rows if is_cd(r) and r["pit_available_to"] == NULL], "ZONE_LINEAGE_MASTER")
    n_ft = write_copy_sheet(wb, FT_SHEET, FT_COLUMNS, latest_versions(rows), CARD_SHEET)
    wb.save(XLSX_PATH)

    print(f"{'record_id':<13}{'zone_code':<16}{'unit':<10}{'unit_code':<16}{'type':<16}{'product':<9}conf")
    for r in rows:
        print(f"{r['record_id']:<13}{r['zone_code']:<16}{r['production_unit_name']:<10}{r['production_unit_code']:<16}"
              f"{r['production_unit_type']:<16}{r['primary_product_association']:<9}{r['zone_type_confidence']}")
    print()
    for l in lineage:
        print(f"{l['lineage_id']:<13}{l['relationship_type']:<10}{l['effective_at']}  {l['_pair']}")
    print(f"\nWrote {len(rows)} zones, {len(lineage)} lineage rows, {n_card} {CARD_SHEET} rows and {n_ft} {FT_SHEET} "
          f"rows to {XLSX_PATH}")
    for n in notes + lineage_notes:
        print("WARNING:", n)

    # Independent check of the written workbook (shares no parsing code with this script).
    print("\nValidating...", flush=True)
    validator = Path(__file__).with_name("validate_zone_master.py")
    sys.exit(subprocess.run([sys.executable, str(validator), str(kml_path)]).returncode)


if __name__ == "__main__":
    main(sys.argv[1])
